from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from chromadb.utils.embedding_functions import DefaultEmbeddingFunction

from src.config import DATA_DIR, CHROMA_DIR, CHUNK_SIZE, CHUNK_OVERLAP


COLLECTION_NAME = "college_academic_documents"


class ChromaEmbeddingAdapter:
    """
    Adapter that allows LangChain Chroma to use Chroma's
    built-in ONNX embedding function.
    """

    def __init__(self):
        self.embedding_function = DefaultEmbeddingFunction()

    def embed_documents(self, texts):
        return self.embedding_function(texts)

    def embed_query(self, text):
        return self.embedding_function([text])[0]


def load_pdf(file_path: Path):
    """Load text from a PDF file."""
    documents = []

    reader = PdfReader(str(file_path))

    for page_number, page in enumerate(reader.pages):
        text = page.extract_text() or ""

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source_file": file_path.name,
                        "file_type": "pdf",
                        "page": page_number,
                    },
                )
            )

    return documents


def load_docx(file_path: Path):
    """Load text from a DOCX file."""
    docx = DocxDocument(str(file_path))

    paragraphs = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also include table contents because college guideline documents
    # often contain schedules, evaluation criteria, and other information
    # inside tables.
    for table in docx.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                paragraphs.append(row_text)

    full_text = "\n".join(paragraphs)

    if not full_text.strip():
        return []

    return [
        Document(
            page_content=full_text,
            metadata={
                "source_file": file_path.name,
                "file_type": "docx",
                "page": None,
            },
        )
    ]


def load_documents():
    """
    Load all supported documents from the data directory.

    Supported:
    - PDF
    - DOCX
    """

    documents = []

    supported_files = sorted(
        [
            file
            for file in DATA_DIR.iterdir()
            if file.is_file()
            and file.suffix.lower() in {".pdf", ".docx"}
        ]
    )

    if not supported_files:
        raise FileNotFoundError(
            f"No PDF or DOCX files found in {DATA_DIR}"
        )

    for file_path in supported_files:
        print(f"Loading: {file_path.name}")

        if file_path.suffix.lower() == ".pdf":
            documents.extend(load_pdf(file_path))

        elif file_path.suffix.lower() == ".docx":
            documents.extend(load_docx(file_path))

    return documents


def split_documents(documents):
    """Split loaded documents into RAG chunks."""

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=[
            "\n\n",
            "\n",
            ". ",
            " ",
            "",
        ],
    )

    chunks = splitter.split_documents(documents)

    return chunks


def get_vector_store():
    """Create/open the Chroma collection."""

    embedding_function = ChromaEmbeddingAdapter()

    return Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=embedding_function,
        persist_directory=str(CHROMA_DIR),
    )


def clear_existing_collection():
    """
    Delete the existing collection so the knowledge base can be
    rebuilt cleanly from all current documents.
    """

    import chromadb

    client = chromadb.PersistentClient(
        path=str(CHROMA_DIR)
    )

    try:
        client.delete_collection(COLLECTION_NAME)
        print(f"Deleted existing Chroma collection: {COLLECTION_NAME}")
    except Exception:
        print("No existing Chroma collection found. Starting fresh.")


def ingest_documents():
    """Rebuild the complete Chroma knowledge base."""

    print("=" * 60)
    print("COLLEGE DOCUMENT INGESTION")
    print("=" * 60)

    print(f"\nData directory: {DATA_DIR}")

    documents = load_documents()

    print(f"\nLoaded document sections: {len(documents)}")

    chunks = split_documents(documents)

    print(f"Created chunks: {len(chunks)}")

    # Rebuild from scratch so existing chunks are not duplicated.
    clear_existing_collection()

    vector_store = get_vector_store()

    print("\nAdding documents to Chroma...")

    vector_store.add_documents(chunks)

    print("\n" + "=" * 60)
    print("INGESTION COMPLETE")
    print("=" * 60)

    print(f"Documents/sections loaded: {len(documents)}")
    print(f"Chunks stored: {len(chunks)}")
    print(f"Chroma directory: {CHROMA_DIR}")

    print("\nSources included:")

    source_counts = {}

    for chunk in chunks:
        source = chunk.metadata.get(
            "source_file",
            "Unknown",
        )

        source_counts[source] = (
            source_counts.get(source, 0) + 1
        )

    for source, count in source_counts.items():
        print(f"- {source}: {count} chunks")


if __name__ == "__main__":
    ingest_documents()